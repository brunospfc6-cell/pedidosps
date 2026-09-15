from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .config import ROOT
from . import db

UPLOAD_DIR = ROOT / "data" / "uploads" / "price-lists"


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s or "").lower())


HEADER_MAP = {
    "materialnumber": "sku",
    "sku": "sku",
    "partnumber": "sku",
    "autodeskmaterialdescription": "name",
    "materialdescription": "name",
    "description": "name",
    "autodesksapmaterialdescription": "sap",
    "sapmaterialdescription": "sap",
    "sap": "sap",
    "productlinedescription": "line",
    "productline": "line",
    "productgroup": "group",
    "usagetype": "usage",
    "licensetype": "license",
    "contractterm": "term",
    "deployment": "deploy",
    "preciorellerusdiva": "price",
    "precioresellerusdiva": "price",
    "precioresellerusd": "price",
    "listpriceusd": "price",
    "listprice": "price",
    "price": "price",
    "notes": "notes",
    "promoname": "promo",
}


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    return str(v).strip()


def _price(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("R$", "").replace("US$", "").replace("$", "")
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    try:
        return float(re.sub(r"[^0-9.\-]", "", s) or 0)
    except ValueError:
        return 0.0


def parse_effective(text: str) -> tuple[str, str]:
    if not text:
        return "", ""
    m = re.search(r"effective[:\s]+(.+?)(?:\s*[-–]\s*|\s+to\s+)(.+)", text, re.I)
    if not m:
        return "", ""
    return m.group(1).strip(" ."), m.group(2).strip(" .")


def label_from_filename(name: str) -> str:
    stem = Path(name).stem.replace("_", " ").replace("-", " ").strip()
    months = {
        "janeiro": "Janeiro", "fevereiro": "Fevereiro", "marco": "Março", "março": "Março",
        "abril": "Abril", "maio": "Maio", "junho": "Junho", "julho": "Julho",
        "agosto": "Agosto", "setembro": "Setembro", "outubro": "Outubro",
        "novembro": "Novembro", "dezembro": "Dezembro",
    }
    parts = re.split(r"[\s.]+", stem)
    if parts and parts[0].lower() in months:
        year = parts[1] if len(parts) > 1 else ""
        return f"{months[parts[0].lower()]} {year}".strip()
    return re.sub(r"\s+", " ", stem) or "Tabela Autodesk"


def parse_json_bytes(raw: bytes) -> dict:
    data = json.loads(raw.decode("utf-8-sig"))
    items = []
    for it in data.get("items") or []:
        sku = (it.get("sku") or "").strip()
        if not sku:
            continue
        items.append(
            {
                "sku": sku,
                "name": (it.get("name") or sku).strip(),
                "sap": it.get("sap") or "",
                "line": it.get("line") or "",
                "group": it.get("group") or "Autodesk",
                "usage": it.get("usage") or "Commercial",
                "term": it.get("term") or "",
                "license": it.get("license") or "",
                "deploy": it.get("deploy") or "",
                "notes": it.get("notes") or "",
                "promo": it.get("promo") or "",
                "price": _price(it.get("price")),
            }
        )
    return {
        "label": data.get("label") or "Tabela Autodesk",
        "version": data.get("version") or datetime.now().date().isoformat(),
        "effectiveFrom": data.get("effectiveFrom") or "",
        "effectiveTo": data.get("effectiveTo") or "",
        "items": items,
    }


def parse_xlsx_bytes(raw: bytes) -> dict:
    wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    sheet = None
    for name in wb.sheetnames:
        if _norm(name) in ("pricelist", "lista", "precos", "preços"):
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    mapping: dict[int, str] = {}
    header_found = False
    effective_from = ""
    effective_to = ""
    items: list[dict] = []

    for row in sheet.iter_rows(values_only=True):
        vals = list(row)
        if not effective_from:
            for cell in vals:
                t = _cell(cell)
                if t.lower().startswith("effective"):
                    effective_from, effective_to = parse_effective(t)
                    break

        if not header_found:
            found: dict[str, int] = {}
            for j, cell in enumerate(vals):
                key = HEADER_MAP.get(_norm(cell))
                if key and key not in found:
                    found[key] = j
            if "sku" in found and ("name" in found or "price" in found):
                header_found = True
                mapping = {j: k for k, j in found.items()}
            continue

        rec = {key: (vals[j] if j < len(vals) else None) for j, key in mapping.items()}
        sku = _cell(rec.get("sku"))
        if not sku:
            continue
        notes = " · ".join(
            x for x in [_cell(rec.get("notes")), _cell(rec.get("promo"))] if x and x not in ("N/A",)
        )
        items.append(
            {
                "sku": sku,
                "name": _cell(rec.get("name")) or sku,
                "sap": _cell(rec.get("sap")),
                "line": _cell(rec.get("line")),
                "group": _cell(rec.get("group")) or "Autodesk",
                "usage": _cell(rec.get("usage")) or "Commercial",
                "term": _cell(rec.get("term")),
                "license": _cell(rec.get("license")),
                "deploy": _cell(rec.get("deploy")),
                "notes": notes,
                "promo": _cell(rec.get("promo")),
                "price": _price(rec.get("price")),
            }
        )
    wb.close()
    if not items:
        raise ValueError(
            "Não encontrei SKUs na planilha. Use a lista Autodesk (aba Price List, coluna Material Number)."
        )
    return {
        "label": "",
        "version": datetime.now().date().isoformat(),
        "effectiveFrom": effective_from,
        "effectiveTo": effective_to,
        "items": items,
    }


def parse_upload(filename: str, raw: bytes) -> dict:
    name = (filename or "tabela.xlsx").lower()
    if name.endswith(".json"):
        data = parse_json_bytes(raw)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        data = parse_xlsx_bytes(raw)
    else:
        raise ValueError("Envie um arquivo .xlsx (planilha Autodesk) ou .json.")
    if not data.get("label"):
        data["label"] = label_from_filename(filename)
    return data


def get_meta(conn=None) -> dict:
    conn = conn or db.get_conn()
    try:
        rows = {r["key"]: r["value"] for r in db.rows(conn.execute("SELECT key, value FROM catalog_meta"))}
    except Exception:
        rows = {}
    count = conn.execute("SELECT COUNT(*) AS n FROM products WHERE active=1").fetchone()["n"]
    return {
        "label": rows.get("price_list_label") or "Tabela Autodesk",
        "version": rows.get("price_list_version") or "",
        "file": rows.get("price_list_file") or "",
        "effective": rows.get("price_list_effective") or "",
        "imported_at": rows.get("price_list_imported_at") or "",
        "count": int(rows.get("price_list_count") or count),
        "active": count,
    }


def set_meta(conn, **pairs):
    for k, v in pairs.items():
        conn.execute(
            "INSERT INTO catalog_meta (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (k, str(v if v is not None else "")),
        )


def save_upload_and_apply(filename: str, raw: bytes, user_id: int) -> dict:
    parsed = parse_upload(filename, raw)
    items = parsed["items"]
    skus = [it["sku"] for it in items]
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", filename or "tabela.xlsx")
    (UPLOAD_DIR / f"{stamp}-{safe}").write_bytes(raw)

    with db.tx() as conn:
        conn.executemany(
            """INSERT INTO products (
                 sku, name, category, sap_name, product_line, usage_type, license_type,
                 contract_term, deployment, notes, list_price_usd, active, source, updated_at
               ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?, 'price_list', datetime('now'))
               ON CONFLICT(sku) DO UPDATE SET
                 name=excluded.name,
                 category=excluded.category,
                 sap_name=excluded.sap_name,
                 product_line=excluded.product_line,
                 usage_type=excluded.usage_type,
                 license_type=excluded.license_type,
                 contract_term=excluded.contract_term,
                 deployment=excluded.deployment,
                 notes=excluded.notes,
                 list_price_usd=excluded.list_price_usd,
                 active=1,
                 source='price_list',
                 updated_at=datetime('now')""",
            [
                (
                    it["sku"],
                    it["name"],
                    it.get("group") or "Autodesk",
                    it.get("sap") or "",
                    it.get("line") or "",
                    it.get("usage") or "Commercial",
                    it.get("license") or "",
                    it.get("term") or "",
                    it.get("deploy") or "",
                    it.get("notes") or "",
                    float(it.get("price") or 0),
                    1,
                )
                for it in items
            ],
        )
        qmarks = ",".join("?" * len(skus))
        conn.execute(
            f"UPDATE products SET active=0, updated_at=datetime('now') WHERE source='price_list' AND sku NOT IN ({qmarks})",
            skus,
        )
        effective = " — ".join(x for x in [parsed.get("effectiveFrom"), parsed.get("effectiveTo")] if x)
        label = parsed.get("label") or label_from_filename(filename)
        set_meta(
            conn,
            price_list_label=label,
            price_list_version=parsed.get("version") or datetime.now().date().isoformat(),
            price_list_file=filename,
            price_list_count=len(items),
            price_list_effective=effective,
            price_list_imported_at=datetime.now().isoformat(timespec="seconds"),
        )
        conn.execute(
            """INSERT INTO price_list_imports (filename, label, sku_count, imported_by, notes)
               VALUES (?,?,?,?,?)""",
            (filename, label, len(items), user_id, effective),
        )
        seed_path = ROOT / "seed" / "price-list.json"
        seed_path.parent.mkdir(parents=True, exist_ok=True)
        seed_path.write_text(
            json.dumps(
                {
                    "version": parsed.get("version"),
                    "effectiveFrom": parsed.get("effectiveFrom"),
                    "effectiveTo": parsed.get("effectiveTo"),
                    "label": label,
                    "items": items,
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
    meta = get_meta()
    meta["imported"] = len(items)
    return meta


def list_imports(limit: int = 12) -> list[dict]:
    return db.rows(
        db.get_conn().execute(
            """SELECT i.*, u.name AS imported_by_name
               FROM price_list_imports i
               LEFT JOIN users u ON u.id = i.imported_by
               ORDER BY i.id DESC LIMIT ?""",
            (limit,),
        )
    )
